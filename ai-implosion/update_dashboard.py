"""AI Implosion Dashboard — daily updater (by KK Reddy).

Pulls ~2 years of closes from Yahoo Finance + HY OAS from FRED, then writes
  data/ai-implosion.json                         (powers kkreckons.com/ai-implosion)
  public/ai-implosion/AI_Implosion_Dashboard.xlsx (the Excel workbook, downloadable from the page)
Weights, Calm/Trigger thresholds and bet picks live in ai-implosion/config.json.
A GitHub Action runs this every weekday after the US close.
Usage:  python ai-implosion/update_dashboard.py [--sample]   (needs: pip install openpyxl)
"""
import datetime as dt, random, math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.table import Table, TableStyleInfo, TableColumn, TableFormula
from openpyxl.formatting.rule import CellIsRule, FormulaRule, DataBarRule
from openpyxl.chart import BarChart, LineChart, DoughnutChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties

RED, ORANGE, GREEN, INDIGO, AMBER, TEAL, SLATE = "EF4444", "D97706", "16A34A", "6366F1", "F59E0B", "14B8A6", "1E293B"
ORANGERED = "EA580C"
F = "Arial"
def font(**k): return Font(name=F, **k)
def fill(c): return PatternFill("solid", start_color=c, end_color=c)
thin = Side(style="thin", color="CBD5E1")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
INPUT = fill("DBEAFE")

FACT = [  # name, col, source, pillar, w, calm, trig, type
 ("Nasdaq-100","NDX","Yahoo ^NDX","Equity",.18,-.05,-.25,"DD"),
 ("Semis","SOX","Yahoo ^SOX","Equity",.18,-.10,-.40,"DD"),
 ("AI ETF","AIQ","Yahoo AIQ","Equity",.09,-.05,-.30,"DD"),
 ("Oracle","ORCL","Yahoo ORCL","Financing",.09,-.15,-.50,"DD"),
 ("CoreWeave","CRWV","Yahoo CRWV","Financing",.09,-.30,-.60,"DD"),
 ("BDC ETF","BIZD","Yahoo BIZD","Credit",.14,-.05,-.25,"DD"),
 ("Loan ETF","BKLN","Yahoo BKLN","Credit",.13,-.02,-.08,"DD"),
 ("HY OAS (bp)","HY","FRED BAMLH0A0HYM2 x100","Credit",.10,400,600,"LVL"),
]
COLS = [f[1] for f in FACT]
T = "tblHistory"
def r(c): return f"{T}[[#This Row],[{c}]]"
def col(c): return f"{T}[{c}]"

wb = Workbook()
dash = wb.active; dash.title = "Dashboard"
trk = wb.create_sheet("Tracker"); hist = wb.create_sheet("History"); cd = wb.create_sheet("ChartData")
for ws in wb.worksheets: ws.sheet_view.showGridLines = False

# ---------------- DATA ----------------
import sys, os
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import engine
SAMPLE = "--sample" in sys.argv
REPO = os.path.dirname(HERE)
OUT = next((a for a in sys.argv[1:] if a.endswith(".xlsx")), os.path.join(REPO, "public", "ai-implosion", "AI_Implosion_Dashboard.xlsx"))
JSON_OUT = os.path.join(REPO, "data", "ai-implosion.json")
cfg = engine.load_config()
FACT = [(f["name"], f["col"], f["source"], f["pillar"], f["weight"], f["calm"], f["trigger"], f["type"]) for f in cfg["factors"]]
PICKS = [(k, dt.date.fromisoformat(v)) for k, v in cfg["bet"].items()]
dates, series = engine.get_data(cfg, SAMPLE)
n = len(dates)
summary = engine.write_json(dates, series, cfg, JSON_OUT, SAMPLE)
os.makedirs(os.path.dirname(OUT), exist_ok=True)

hist["A1"] = ("SAMPLE DATA — illustrative random-walk prices, NOT real market history. Run update_dashboard.py to load real data." if SAMPLE else
   f"LIVE DATA — Yahoo Finance closes + FRED BAMLH0A0HYM2 (x100). Auto-refreshed {dt.datetime.now():%d-%b-%Y %H:%M}. Newest row on top.")
hist["A1"].font = font(bold=True, color=RED if SAMPLE else GREEN)
hist["A2"] = "Inputs: Date + 8 closes (light blue). Everything from 'NDX S' onward is calculated — don't type there. Blank HY = carry last value."
hist["A2"].font = font(italic=True, color="475569", size=9)

S = [c + " S" for c in COLS]
headers = ["Date"] + COLS + S + ["EquityP","FinP","CreditP","Score","L1","L2","L3","L4","L1S","L2S","L3S","L4S"]
HR = 3
for j, h in enumerate(headers, 1):
    c = hist.cell(HR, j, h); c.font = font(bold=True, color="FFFFFF"); c.fill = fill(SLATE); c.alignment = Alignment(horizontal="center")

wu = f"{r('Date')}-MIN({col('Date')})<364"
hyx = f"IF(ISNUMBER({r('HY')}),{r('HY')},SUMIFS({col('HY')},{col('Date')},_xlfn.MAXIFS({col('Date')},{col('Date')},\"<\"&{r('Date')},{col('HY')},\"<>\")))"
F_ = {}
for k, f in enumerate(FACT):
    dr = 13 + k; c = f[1]
    if f[7] == "DD":
        x = f"({r(c)}/_xlfn.MAXIFS({col(c)},{col('Date')},\"<=\"&{r('Date')},{col('Date')},\">=\"&{r('Date')}-365)-1)"
        F_[c + " S"] = f"IF({wu},NA(),IF(ISNUMBER({r(c)}),MIN(1,MAX(0,({x}-Dashboard!$G${dr})/(Dashboard!$H${dr}-Dashboard!$G${dr}))),0))"
    else:
        F_[c + " S"] = f"IF({wu},NA(),MIN(1,MAX(0,({hyx}-Dashboard!$G${dr})/(Dashboard!$H${dr}-Dashboard!$G${dr}))))"
def wavg(rows):
    num = "+".join(f"Dashboard!$C${13+i}*{r(COLS[i]+' S')}" for i in rows)
    return f"({num})/SUM(Dashboard!$C${13+rows[0]}:$C${13+rows[-1]})*100"
F_["EquityP"] = wavg([0,1,2]); F_["FinP"] = wavg([3,4]); F_["CreditP"] = wavg([5,6,7])
F_["Score"] = "ROUND(" + wavg(list(range(8))).replace("SUM(Dashboard!$C$13:$C$20)", "SUM(Dashboard!$C$13:$C$20)") + ",0)"
F_["L1"] = f"IFERROR(--({r('Score')}>=50),0)"
F_["L2"] = f"IFERROR(--AND({r('Score')}>=60,{r('BIZD S')}>=0.8,{r('BKLN S')}>=0.5),0)"
F_["L3"] = f"IFERROR(--AND({r('Score')}>=70,{r('SOX S')}>=0.83,{r('AIQ S')}>=0.8),0)"
F_["L4"] = f"IFERROR(--AND({r('Score')}>=75,{r('NDX S')}=1,OR({hyx}>=500,{r('BIZD S')}=1)),0)"
for L in ["L1","L2","L3","L4"]:
    F_[L + "S"] = (f"IFERROR(--(COUNTIFS({col('Date')},\">=\"&{r('Date')},{col('Date')},\"<=\"&"
                   f"_xlfn.AGGREGATE(15,6,{col('Date')}/({col('Date')}>={r('Date')}),5),{col(L)},1)=5),0)")

rows_desc = list(range(n - 1, -1, -1))  # newest first
for i, src in enumerate(rows_desc):
    rr = HR + 1 + i
    hist.cell(rr, 1, dates[src]).number_format = "dd-mmm-yyyy"
    for j, c in enumerate(COLS):
        cell = hist.cell(rr, 2 + j, series[c][src]); cell.fill = INPUT
        cell.number_format = "#,##0" if c == "HY" else "#,##0.00"
    for j, h in enumerate(headers[9:], 10):
        cell = hist.cell(rr, j, "=" + F_[h])
        cell.number_format = "0%" if h.endswith(" S") else ("0" if h in ("EquityP","FinP","CreditP","Score") or h.startswith("L") else "General")
last = HR + n
ref = f"A{HR}:{hist.cell(HR, len(headers)).column_letter}{last}"
tab = Table(displayName=T, ref=ref)
tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
tab._initialise_columns()
for tc in tab.tableColumns:
    tc.name = tc.name  # header names
for tc, h in zip(tab.tableColumns, headers):
    tc.name = h
    if h in F_: tc.calculatedColumnFormula = TableFormula(attr_text=F_[h])
hist.add_table(tab)
hist.freeze_panes = "B4"
hist.column_dimensions["A"].width = 13
for j in range(2, len(headers) + 1):
    hist.column_dimensions[hist.cell(HR, j).column_letter].width = 10

# ---------------- DASHBOARD ----------------
dash.column_dimensions["A"].width = 2
for c, w in zip("BCDEFGHIJK", [16, 10, 13, 13, 15, 10, 10, 16, 14, 13]): dash.column_dimensions[c].width = w
dash.merge_cells("B2:K2"); dash["B2"] = "AI IMPLOSION DASHBOARD"
dash["B2"].font = font(bold=True, size=22, color="FFFFFF"); dash["B2"].fill = fill(SLATE)
dash.row_dimensions[2].height = 36; dash["B2"].alignment = Alignment(vertical="center", indent=1)
dash.merge_cells("B3:K3"); dash["B3"] = "Built by KK Reddy"; dash["B3"].font = font(italic=True, size=9, color="64748B")
dash.merge_cells("B4:K4")
dash["B4"] = f'="7 AI-cycle drawdowns + HY spread | Data as of "&TEXT(MAX({col("Date")}),"dd-mmm-yyyy")'
dash["B4"].font = font(size=11, color="334155")

cards = [("B","C","STRESS SCORE","=Tracker!C5","=\"out of 100\""),
         ("D","E","STATUS","=Tracker!C7","=\"Score band\""),
         ("F","G","TRIGGERS","=Tracker!C6","=\"factors at stress 100%\""),
         ("H","I","CURRENT LEVEL","=Tracker!C17","=\"5 straight days rule\""),
         ("J","K","BET COUNTDOWN","=\"KK \"&Tracker!D22&\"d | SJ \"&Tracker!D23&\"d\"","=\"Winner: \"&Tracker!C25")]
for a, b, lab, val, sub in cards:
    for rr in (7, 8, 9):
        dash.merge_cells(f"{a}{rr}:{b}{rr}")
        for cc in (a, b):
            dash[f"{cc}{rr}"].fill = fill("F1F5F9");
    dash[f"{a}7"] = lab; dash[f"{a}7"].font = font(bold=True, size=9, color="64748B")
    dash[f"{a}8"] = val; dash[f"{a}8"].font = font(bold=True, size=18 if a != "J" else 13, color=SLATE)
    dash[f"{a}9"] = sub; dash[f"{a}9"].font = font(size=9, color="64748B")
    for rr in (7, 8, 9): dash[f"{a}{rr}"].alignment = Alignment(horizontal="center", vertical="center")
dash.row_dimensions[8].height = 30
dash["B8"].number_format = "0"
for colr, f_ in [(RED, "$B$8>=75"), (ORANGERED, "$B$8>=50"), (ORANGE, "$B$8>=25"), (GREEN, "ISNUMBER($B$8)")]:
    dash.conditional_formatting.add("B8", FormulaRule(formula=[f_], font=Font(name=F, bold=True, size=18, color=colr), stopIfTrue=True))

hdr = ["Metric","Weight","52W High (HY: 52W Low)","Current","Drawdown (HY: +bp vs low)","Calm","Trigger","Trigger price (HY: band)","Stress","Status"]
dash.merge_cells("B11:K11"); dash["B11"] = "FACTOR DETAIL"; dash["B11"].font = font(bold=True, size=11, color=SLATE)
for j, h in enumerate(hdr):
    c = dash.cell(12, 2 + j, h); c.font = font(bold=True, color="FFFFFF", size=9); c.fill = fill(SLATE)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = box
dash.row_dimensions[12].height = 30
MAXD = f"MAX({col('Date')})"
for k, f in enumerate(FACT):
    rr = 13 + k; c = f[1]; dd = f[7] == "DD"
    cur = f"SUMIFS({col(c)},{col('Date')},_xlfn.MAXIFS({col('Date')},{col(c)},\"<>\"))"
    vals = [f"{f[0]} ({c})", f[4],
            (f"=_xlfn.MAXIFS({col(c)},{col('Date')},\"<=\"&{MAXD},{col('Date')},\">=\"&{MAXD}-365)" if dd else
             f"=_xlfn.MINIFS({col(c)},{col('Date')},\"<=\"&{MAXD},{col('Date')},\">=\"&{MAXD}-365)"),
            "=" + cur,
            f"=E{rr}/D{rr}-1" if dd else f"=E{rr}-D{rr}",
            f[5], f[6],
            f"=D{rr}*(1+H{rr})" if dd else
            f'=IF(E{rr}<300,"Very Healthy",IF(E{rr}<400,"Normal",IF(E{rr}<500,"Watch",IF(E{rr}<600,"Warning",IF(E{rr}<800,"Severe Stress","Crisis")))))',
            (f"=MIN(1,MAX(0,(F{rr}-G{rr})/(H{rr}-G{rr})))" if dd else f"=MIN(1,MAX(0,(E{rr}-G{rr})/(H{rr}-G{rr})))"),
            f'=IF(J{rr}=1,"TRIGGERED",IF(J{rr}>=0.5,"WARNING",IF(J{rr}>0,"WATCH","CALM")))']
    fmts = ["@", "0%", "#,##0.00" if dd else "#,##0", "#,##0.00" if dd else "#,##0",
            "0.0%" if dd else "+#,##0;-#,##0;0", "0%" if dd else "#,##0", "0%" if dd else "#,##0",
            "#,##0.00" if dd else "@", "0%", "@"]
    for j, (v, fm) in enumerate(zip(vals, fmts)):
        cell = dash.cell(rr, 2 + j, v); cell.number_format = fm; cell.border = box
        cell.font = font(size=10); cell.alignment = Alignment(horizontal="left" if j == 0 else "center")
    for cc in ("C", "G", "H"):
        dash[f"{cc}{rr}"].fill = INPUT; dash[f"{cc}{rr}"].protection = Protection(locked=False)
dash.cell(21, 2, "Total weight"); dash["C21"] = "=SUM(C13:C20)"; dash["C21"].number_format = "0%"
dash["B21"].font = font(bold=True, size=9); dash["C21"].font = font(bold=True, size=9)
dash.conditional_formatting.add("K13:K20", CellIsRule(operator="equal", formula=['"TRIGGERED"'], fill=fill(RED), font=Font(name=F, bold=True, color="FFFFFF")))
dash.conditional_formatting.add("K13:K20", CellIsRule(operator="equal", formula=['"WARNING"'], fill=fill("FED7AA")))
dash.conditional_formatting.add("K13:K20", CellIsRule(operator="equal", formula=['"WATCH"'], fill=fill("FEF9C3")))
dash.conditional_formatting.add("K13:K20", CellIsRule(operator="equal", formula=['"CALM"'], fill=fill("DCFCE7")))
dash.conditional_formatting.add("J13:J20", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=INDIGO))
dash.merge_cells("B22:K23")
dash["B22"] = ("How it works: Drawdown = close ÷ trailing-365-day high − 1 (HY uses spread level in bp). "
               "Stress = clamp((x − Calm)/(Trigger − Calm), 0, 1). Score = weighted stress × 100. "
               "Levels need 5 straight trading days: L1 Score≥50 · L2 Score≥60 + BIZD≥80% + BKLN≥50% · "
               "L3 Score≥70 + SOX≥83% + AIQ≥80% · L4 Score≥75 + NDX=100% + (HY≥500bp or BIZD=100%). "
               "Implosion date = first day of the first L4 streak. Blue cells are editable.")
dash["B22"].font = font(size=9, italic=True, color="475569"); dash["B22"].alignment = Alignment(wrap_text=True, vertical="top")
dash.row_dimensions[22].height = 30; dash.row_dimensions[23].height = 30

# ---------------- TRACKER ----------------
trk.column_dimensions["A"].width = 2
for c, w in zip("BCDEF", [24, 22, 16, 14, 14]): trk.column_dimensions[c].width = w
trk["B2"] = "TRACKER — score, levels & bet"; trk["B2"].font = font(bold=True, size=16, color=SLATE)
kv = [("Latest date", f"={MAXD}", "dd-mmm-yyyy"),
      ("Score", f"=INDEX({col('Score')},MATCH(C4,{col('Date')},0))", "0"),
      ("Triggers", '=COUNTIF(Dashboard!K13:K20,"TRIGGERED")&" of 8"', "@"),
      ("Band", '=IF(C18<>"","IMPLOSION CONFIRMED",IFERROR(IF(C5<25,"NORMAL",IF(C5<50,"ELEVATED",IF(C5<75,"SEVERE","CRITICAL"))),"WARM-UP"))', "@")]
for i, (k, v, fm) in enumerate(kv):
    trk.cell(4 + i, 2, k).font = font(bold=True); c = trk.cell(4 + i, 3, v); c.number_format = fm; c.font = font()
trk["B10"] = "LEVELS"; trk["B10"].font = font(bold=True, size=11, color=SLATE)
for j, h in enumerate(["Level", "Name", "First date (5-day start)", "Active now"]):
    c = trk.cell(11, 2 + j, h); c.font = font(bold=True, color="FFFFFF"); c.fill = fill(SLATE); c.border = box
names = ["Warning", "Financing Stress", "Bubble Unwind", "AI Implosion"]
for i in range(4):
    rr = 12 + i; L = f"L{i+1}"
    vals = [L, names[i],
            f'=IF(COUNTIF({col(L+"S")},1),_xlfn.MINIFS({col("Date")},{col(L+"S")},1),"")',
            f'=IF(COUNTIFS({col("Date")},">="&LARGE({col("Date")},5),{col(L)},1)=5,"YES","NO")']
    for j, v in enumerate(vals):
        c = trk.cell(rr, 2 + j, v); c.font = font(); c.border = box
    trk.cell(rr, 4).number_format = "dd-mmm-yyyy"
trk["B17"] = "Current level"; trk["B17"].font = font(bold=True)
trk["C17"] = '=IF(E15="YES","L4 AI Implosion",IF(E14="YES","L3 Bubble Unwind",IF(E13="YES","L2 Financing Stress",IF(E12="YES","L1 Warning","None"))))'
trk["B18"] = "AI Implosion date"; trk["B18"].font = font(bold=True, color=RED)
trk["C18"] = "=D15"; trk["C18"].number_format = "dd-mmm-yyyy"; trk["C18"].font = font(bold=True, color=RED)
trk["B20"] = "THE BET"; trk["B20"].font = font(bold=True, size=11, color=SLATE)
for j, h in enumerate(["Player", "Pick", "Days left", "Days off"]):
    c = trk.cell(21, 2 + j, h); c.font = font(bold=True, color="FFFFFF"); c.fill = fill(SLATE); c.border = box
for i, (who, pick) in enumerate(PICKS):
    rr = 22 + i
    trk.cell(rr, 2, who); c = trk.cell(rr, 3, pick); c.number_format = "dd-mmm-yyyy"; c.fill = INPUT; c.protection = Protection(locked=False)
    trk.cell(rr, 4, f"=MAX(0,C{rr}-TODAY())")
    trk.cell(rr, 5, f'=IF($C$18="","",ABS($C$18-C{rr}))')
    for j in range(4): trk.cell(rr, 2 + j).border = box; trk.cell(rr, 2 + j).font = font()
trk["B25"] = "Winner"; trk["B25"].font = font(bold=True)
trk["C25"] = '=IF(C18="","TBD",IF(E22<E23,"KK",IF(E23<E22,"SJ","TIE")))'; trk["C25"].font = font(bold=True, color=GREEN)

# ---------------- CHARTDATA ----------------
cd["A1"] = "Factor"; cd["B1"] = "Stress"
for k, f in enumerate(FACT):
    cd.cell(2 + k, 1, f[1]); c = cd.cell(2 + k, 2, f"=Dashboard!J{13+k}"); c.number_format = "0%"
cd["D1"] = "Gauge"; cd["E1"] = "Value"
for i, (lab, v) in enumerate([("Score", "=IFERROR(Tracker!C5,0)"), ("Rest", "=100-E2"), ("Hidden", 100)]):
    cd.cell(2 + i, 4, lab); cd.cell(2 + i, 5, v)
cd["G1"] = "Week ending (Fri)"; cd["H1"] = "Weekly high score"
for i in range(26):
    rr = 2 + i
    cd.cell(rr, 7, f"={MAXD}-WEEKDAY({MAXD},2)+5" if i == 0 else f"=G{rr-1}-7").number_format = "dd-mmm"
    cd.cell(rr, 8, f'=IF(COUNTIFS({col("Date")},">="&G{rr}-6,{col("Date")},"<="&G{rr})=0,NA(),_xlfn.MAXIFS({col("Score")},{col("Date")},">="&G{rr}-6,{col("Date")},"<="&G{rr}))')
for j, h in enumerate(["Date", "EquityP", "FinP", "CreditP", "Score"]): cd.cell(1, 10 + j, h)
NLINE = 262
for i in range(NLINE):
    rr = 2 + i; nidx = i + 1
    inwin = f"INDEX({col('Date')},{nidx})>{MAXD}-364"
    cd.cell(rr, 10, f'=IFERROR(IF({inwin},INDEX({col("Date")},{nidx}),""),"")').number_format = "dd-mmm-yy"
    for j, h in enumerate(["EquityP", "FinP", "CreditP", "Score"]):
        cd.cell(rr, 11 + j, f'=IF(J{rr}="",NA(),IFERROR(INDEX({col(h)},{nidx}),NA()))').number_format = "0"
for ws in (cd,):
    for c in "ABDEGHJKLMN": ws.column_dimensions[c].width = 12
    ws["A1"].font = font(bold=True)

# ---------------- CHARTS ----------------
def axes_on(ch):
    ch.x_axis.delete = False; ch.y_axis.delete = False

# 1 horizontal bar
b = BarChart(); b.type = "bar"; b.style = 10; b.title = "Factor stress"
b.add_data(Reference(cd, min_col=2, min_row=1, max_row=9), titles_from_data=True)
b.set_categories(Reference(cd, min_col=1, min_row=2, max_row=9))
b.y_axis.scaling.min = 0; b.y_axis.scaling.max = 1; b.y_axis.number_format = "0%"; b.y_axis.majorGridlines = None
b.x_axis.scaling.orientation = "maxMin"
b.legend = None; b.series[0].graphicalProperties.solidFill = INDIGO
b.dataLabels = DataLabelList(); b.dataLabels.showVal = True; b.dataLabels.numFmt = "0%"
axes_on(b); b.width = 13; b.height = 8
dash.add_chart(b, "B25")

# 2 gauge
g = DoughnutChart(); g.title = "Stress Score (0–100)"; g.firstSliceAng = 270; g.holeSize = 60
g.add_data(Reference(cd, min_col=5, min_row=2, max_row=4))
g.set_categories(Reference(cd, min_col=4, min_row=2, max_row=4))
s = g.series[0]
for idx, clr in enumerate([RED, "E2E8F0", None]):
    pt = DataPoint(idx=idx)
    if clr: pt.graphicalProperties.solidFill = clr
    else: pt.graphicalProperties.noFill = True
    pt.graphicalProperties.line.noFill = True
    s.dPt.append(pt)
g.legend = None; g.width = 11; g.height = 8
dash.add_chart(g, "H25")

# 3 lines
lc = LineChart(); lc.title = "Pillars & Score — last 52 weeks"; lc.style = 2
lc.add_data(Reference(cd, min_col=11, max_col=14, min_row=1, max_row=1 + NLINE), titles_from_data=True)
lc.set_categories(Reference(cd, min_col=10, min_row=2, max_row=1 + NLINE))
for ser, clr, wdt in zip(lc.series, [INDIGO, AMBER, TEAL, RED], [15000, 15000, 15000, 28000]):
    ser.graphicalProperties.line.solidFill = clr; ser.graphicalProperties.line.width = wdt; ser.smooth = False
lc.y_axis.scaling.min = 0; lc.y_axis.scaling.max = 100; lc.y_axis.majorUnit = 25
lc.x_axis.scaling.orientation = "maxMin"; lc.y_axis.crosses = "max"
lc.x_axis.number_format = "mmm-yy"; lc.x_axis.tickLblSkip = 21; lc.x_axis.tickMarkSkip = 21
lc.legend.position = "b"; axes_on(lc); lc.width = 24; lc.height = 8.5
dash.add_chart(lc, "B42")

# 4 weekly high
w = BarChart(); w.type = "col"; w.title = "Weekly high score — 26 weeks"
w.add_data(Reference(cd, min_col=8, min_row=1, max_row=27), titles_from_data=True)
w.set_categories(Reference(cd, min_col=7, min_row=2, max_row=27))
w.y_axis.scaling.min = 0; w.y_axis.scaling.max = 100; w.y_axis.majorUnit = 25
w.x_axis.scaling.orientation = "maxMin"; w.y_axis.crosses = "max"; w.x_axis.number_format = "dd-mmm"
w.legend = None; w.series[0].graphicalProperties.solidFill = ORANGE; axes_on(w); w.width = 24; w.height = 8
dash.add_chart(w, "B60")
dash.merge_cells("B77:K77"); dash["B77"] = "AI Implosion Dashboard · designed by KK Reddy"
dash["B77"].font = font(italic=True, size=9, color="64748B"); dash["B77"].alignment = Alignment(horizontal="center")

# ---------------- PROTECTION ----------------
for ws in (dash, trk, cd):
    ws.protection.sheet = True
    ws.protection.formatCells = False; ws.protection.formatColumns = False; ws.protection.formatRows = False
wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print(f"Updated {OUT} and {JSON_OUT}: {n} rows, latest {dates[-1]}, score {summary['score']} ({summary['band']})")
